#!/usr/bin/env python3
"""Both calculators agree on filler/input_week_test.txt.

Hand totals use roundtable entry 006. Full session is 30 in this fixture
because every column C value is "30 min". The filler still reads column C;
that base is not changed until the 45-minute question is answered.

3/2 codes: T 30 + G1 once 30 + T/ 15 + TT/ (30+15) 45 = 120. Note (30 min) = 30. Total 150.
3/3 codes: TD (30+30) 60 + G2 once 30 + G1/ 15 + A 0 + A1 0 = 105. Note (1hr) = 60. Total 165.
3/4 codes: A2 0 + O 0 + H 0 + S 30 + S/ 15 + C 30 = 75. Note (1hr 15 min) = 75. Total 150.
3/5 codes: CC/ (30+15) 45 + R/C (15+15) 30 + D// billed as 15 = 90. Note (45 mins) = 45. Total 135.
3/6 codes: T 30. No note. Total 30.
"""

import io
import subprocess
import sys
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

import smr_filler

FIXTURE = Path(__file__).with_name("input_week_test.txt")
APP = Path(__file__).resolve().parents[1] / "smr-app.html"

# date -> total minutes, codes plus notes
EXPECTED = {
    "3/2": 150,
    "3/3": 165,
    "3/4": 150,
    "3/5": 135,
    "3/6": 30,
}
# date -> code minutes only, before notes
EXPECTED_CODES = {
    "3/2": 120,
    "3/3": 105,
    "3/4": 75,
    "3/5": 90,
    "3/6": 30,
}

STUDENTS = [
    "Alvarez Mia",
    "Alvarez Leo",
    "Brooks Jonah (A)",
    "Brooks Jonah (B)",
    "Chen Priya",
    "Diaz Omar",
]
DATES = [(3, 2), (3, 3), (3, 4), (3, 5), (3, 6)]


def build_sheet(sessions):
    wb = Workbook()
    ws = wb.active
    ws.title = smr_filler.SHEET_NAME
    date_map = {}
    for index, (month, day) in enumerate(DATES):
        col = smr_filler.DATE_COL_START + index
        ws.cell(row=5, column=col, value=datetime(2026, month, day))
        date_map[(month, day)] = col
    rows = {}
    for index, name in enumerate(STUDENTS):
        row = smr_filler.STUDENT_ROW_START + index
        rows[name] = row
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=3, value="30 min")
    for date_str, name, code in sessions:
        month, day = (int(part) for part in date_str.split("/"))
        ws.cell(row=rows[name], column=date_map[(month, day)], value=code)
    return ws, date_map


def notes_text(notes):
    lines = []
    for date_str, text in notes:
        lines.append(date_str)
        lines.append(text)
    return "\n".join(lines)


def filler_totals():
    sessions, _comments, notes = smr_filler.parse_input(str(FIXTURE))
    ws, date_map = build_sheet(sessions)
    buf = io.StringIO()
    with redirect_stdout(buf):
        totals = smr_filler.calculate_daily_totals(ws, ws, date_map, notes_text(notes))
    return totals, buf.getvalue(), sessions


def app_code_totals():
    html = APP.read_text()
    note_start = html.index("function parseNoteMins")
    note_end = html.index("function dateTotalMins")
    start = html.index("// Billing numbers and code classes.")
    end = html.index("// ── Counts ──")
    section = html[note_start:note_end] + "\n" + html[start:end]
    script = r"""
function fmtDateKey(d) {
  return d.getFullYear() + '-' + String(d.getMonth()+1).padStart(2,'0') + '-' + String(d.getDate()).padStart(2,'0');
}
function sessionKey(studentName, date) {
  return studentName + '|' + fmtDateKey(date);
}
let appData = {students: [], sessions: {}};
""" + section + r"""
const days = [
  ['3/2', new Date(2026, 2, 2)],
  ['3/3', new Date(2026, 2, 3)],
  ['3/4', new Date(2026, 2, 4)],
  ['3/5', new Date(2026, 2, 5)],
  ['3/6', new Date(2026, 2, 6)],
];
const rows = [
  ['3/2', 'Alvarez Mia', 'T'],
  ['3/2', 'Alvarez Leo', 'G1'],
  ['3/2', 'Brooks Jonah (A)', 'G1'],
  ['3/2', 'Brooks Jonah (B)', 'G1'],
  ['3/2', 'Chen Priya', 'T/'],
  ['3/2', 'Diaz Omar', 'TT/'],
  ['3/3', 'Alvarez Mia', 'TD'],
  ['3/3', 'Alvarez Leo', 'G2'],
  ['3/3', 'Brooks Jonah (A)', 'G2'],
  ['3/3', 'Brooks Jonah (B)', 'G1/'],
  ['3/3', 'Chen Priya', 'A'],
  ['3/3', 'Diaz Omar', 'A1'],
  ['3/4', 'Alvarez Mia', 'A2'],
  ['3/4', 'Alvarez Leo', 'O'],
  ['3/4', 'Brooks Jonah (A)', 'H'],
  ['3/4', 'Brooks Jonah (B)', 'S'],
  ['3/4', 'Chen Priya', 'S/'],
  ['3/4', 'Diaz Omar', 'C'],
  ['3/5', 'Alvarez Mia', 'CC/'],
  ['3/5', 'Alvarez Leo', 'R/C'],
  ['3/5', 'Brooks Jonah (A)', 'D//'],
  ['3/6', 'Diaz Omar', 'T'],
];
const names = [...new Set(rows.map(r => r[1]))];
appData.students = names.map(name => ({name}));
const out = {};
for (const [label, date] of days) {
  appData.sessions = {};
  for (const [d, name, code] of rows) {
    if (d === label) appData.sessions[sessionKey(name, date)] = code;
  }
  const tot = calcDayCodeTotals(date);
  out[label] = tot.direct + tot.indirect;
}
const rc = tokenizeCode('R/C');
const dbl = tokenizeCode('D//');
const notes = {
  a: parseNoteMins('30 min'),
  b: parseNoteMins('1hr'),
  c: parseNoteMins('1hr 15 min'),
  d: parseNoteMins('45 mins'),
};
console.log(JSON.stringify({codes: out, rc, dbl, notes}));
"""
    result = subprocess.run(["node", "-e", script], capture_output=True, text=True)
    if result.returncode != 0:
        sys.stderr.write(result.stderr)
        raise SystemExit(result.returncode)
    import json
    return json.loads(result.stdout)


def main():
    totals, warnings, sessions = filler_totals()
    assert smr_filler.parse_code_cell("R/C") == [("R", 0.5), ("C", 0.5)]
    assert smr_filler.parse_code_cell("D//") == [("D", 0.25)]
    assert smr_filler.parse_code_cell("A1") == [("A1", 1.0)]
    assert smr_filler.parse_code_cell("TT/") == [("T", 1.0), ("T", 0.5)]
    assert smr_filler.parse_code_cell("S") == [("S", 1.0)]
    assert smr_filler.parse_code_cell("S/") == [("S", 0.5)]
    assert "row " in warnings and "column " in warnings
    for key, expected in EXPECTED.items():
        assert totals[key] == expected, (key, totals[key], expected)
    app = app_code_totals()
    for key, expected in EXPECTED_CODES.items():
        assert app["codes"][key] == expected, (key, app["codes"][key], expected)
    assert app["rc"][0]["modifier"] == "/" and app["rc"][1]["modifier"] == "/"
    assert app["dbl"][0]["invalid"] == "//"
    assert app["dbl"][0]["modifier"] == "/"
    assert app["notes"] == {"a": 30, "b": 60, "c": 75, "d": 45}
    assert len(sessions) == 22
    print("parity ok", totals)


if __name__ == "__main__":
    main()
